from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import database


def build_export():
    rows = database.all_shifts_for_export()

    wb = Workbook()
    ws = wb.active
    ws.title = "Shifts"

    headers = ["Date", "Worker", "Location", "Museum area"]
    header_fill = PatternFill(start_color="1F3A2E", end_color="1F3A2E", fill_type="solid")
    header_font = Font(color="F1E9D8", bold=True)

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left")

    for row_idx, r in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=r["date"])
        ws.cell(row=row_idx, column=2, value=r["worker"])
        ws.cell(row=row_idx, column=3, value=r["location"])
        ws.cell(row=row_idx, column=4, value=r["area"])

    for col in range(1, len(headers) + 1):
        key = list(rows[0].keys())[col - 1] if rows else None
        max_len = max([len(str(headers[col - 1]))] + [len(str(r[key])) for r in rows]) if rows else len(str(headers[col - 1]))
        ws.column_dimensions[get_column_letter(col)].width = max_len + 4

    # Summary sheet: per-worker totals
    ws2 = wb.create_sheet("Worker Summary")
    ws2.append(["Worker", "Total Shifts"])
    for col in range(1, 3):
        ws2.cell(row=1, column=col).fill = header_fill
        ws2.cell(row=1, column=col).font = header_font

    stats = database.overview_stats()
    for w in sorted(stats["workers"], key=lambda x: -x["shift_count"]):
        ws2.append([w["name"], w["shift_count"]])
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 14

    # Summary sheet: per-location totals (with museum area)
    ws3 = wb.create_sheet("Location Summary")
    ws3.append(["Location", "Museum area", "Total Shifts"])
    for col in range(1, 4):
        ws3.cell(row=1, column=col).fill = header_fill
        ws3.cell(row=1, column=col).font = header_font
    for l in sorted(stats["locations"], key=lambda x: -x["shift_count"]):
        ws3.append([l["name"], l["group_name"], l["shift_count"]])
    ws3.column_dimensions["A"].width = 24
    ws3.column_dimensions["B"].width = 14
    ws3.column_dimensions["C"].width = 14

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
